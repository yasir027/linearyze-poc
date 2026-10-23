"""Read text out of uploaded files. Each section carries a location so answers can point back to it.

Scope: text layers only. Scanned pages, password-protected files and hidden sheets are not read.
"""
import io
from dataclasses import dataclass


class UnsupportedFileType(ValueError):
    pass


@dataclass(frozen=True)
class Section:
    location: str
    text: str


def extract_sections(filename: str, content: bytes) -> list[Section]:
    name = filename.lower()
    if name.endswith((".txt", ".md")):
        return [Section("full text", content.decode("utf-8", errors="replace"))]

    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return [
            Section(f"page {number}", page.extract_text() or "")
            for number, page in enumerate(reader.pages, start=1)
        ]

    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(content))
        sections: list[Section] = []
        heading = "start of document"
        buffer: list[str] = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style.name.lower().startswith("heading"):
                if buffer:
                    sections.append(Section(heading, "\n".join(buffer)))
                    buffer = []
                heading = text
            else:
                buffer.append(text)
        if buffer:
            sections.append(Section(heading, "\n".join(buffer)))
        return sections

    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sections = []
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    sections.append(Section(f"sheet {sheet.title}, row {row_number}", " | ".join(cells)))
        return sections

    raise UnsupportedFileType(f"Unsupported file type: {filename}")
